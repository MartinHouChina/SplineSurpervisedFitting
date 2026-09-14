from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, "/tmp/houcode_pptx_lib")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = ROOT / "outputs" / "presentations"
ASSETS = OUT / "four_datasets_assets"
PPTX_PATH = OUT / "四个数据集介绍_单页.pptx"

NAVY = RGBColor(16, 30, 54)
MUTED = RGBColor(91, 105, 125)
PAPER = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(19, 145, 140)
BLUE = RGBColor(48, 104, 189)
ORANGE = RGBColor(225, 128, 45)
PURPLE = RGBColor(124, 82, 173)


def add_text(slide, x, y, w, h, text, size, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, font="Noto Sans CJK SC"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def curve_score(points: np.ndarray) -> float:
    points = np.asarray(points, dtype=float)
    if len(points) < 8:
        return -1.0
    delta = np.diff(points[:, :2], axis=0)
    lengths = np.linalg.norm(delta, axis=1)
    keep = lengths > 1e-12
    delta = delta[keep]
    if len(delta) < 5:
        return -1.0
    angles = np.unwrap(np.arctan2(delta[:, 1], delta[:, 0]))
    return float(np.quantile(np.abs(np.diff(angles)), 0.9) + 0.02 * np.std(angles))


def representative_from_manifest(manifest: Path, limit=500) -> np.ndarray:
    entries = []
    with manifest.open(encoding="utf-8") as stream:
        for line in stream:
            row = json.loads(line)
            if row.get("split") == "test":
                entries.append(row)
            if len(entries) >= limit:
                break
    if not entries:
        with manifest.open(encoding="utf-8") as stream:
            entries = [json.loads(line) for line in stream][:limit]
    candidates = []
    for row in entries:
        path = (manifest.parent / row["points_path"]).resolve()
        if path.exists():
            pts = np.load(path)
            candidates.append((curve_score(pts), pts))
    if not candidates:
        raise FileNotFoundError(f"No curve arrays resolved from {manifest}")
    return max(candidates, key=lambda item: item[0])[1]


def synthetic_curve() -> np.ndarray:
    t = np.linspace(0, 1, 300)
    x = t + 0.035 * np.sin(10 * np.pi * t)
    y = 0.48 * np.sin(2 * np.pi * t) + 0.15 * np.sin(7 * np.pi * t + 0.4)
    return np.column_stack([x, y])


def save_curve(points: np.ndarray, path: Path, color: str):
    pts = np.asarray(points, dtype=float)[:, :2]
    center = pts.mean(axis=0)
    scale = max(np.ptp(pts[:, 0]), np.ptp(pts[:, 1]), 1e-9)
    pts = (pts - center) / scale
    fig, ax = plt.subplots(figsize=(4.2, 2.0), dpi=180)
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")
    ax.plot(pts[:, 0], pts[:, 1], color=color, linewidth=3.0, solid_capstyle="round")
    ids = np.linspace(0, len(pts) - 1, min(22, len(pts)), dtype=int)
    ax.scatter(pts[ids, 0], pts[ids, 1], s=9, color="white", edgecolor=color,
               linewidth=0.8, zorder=3)
    ax.set_aspect("equal", adjustable="datalim")
    ax.axis("off")
    fig.subplots_adjust(0.02, 0.04, 0.98, 0.96)
    fig.savefig(path, transparent=True, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


def add_card(slide, x, y, title, tag, bullets, image_path, accent):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                  Inches(6.05), Inches(2.52))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(225, 231, 239)
    card.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(0.11), Inches(2.52))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, x + 0.30, y + 0.18, 3.25, 0.38, title, 18, NAVY, True)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 4.62),
                                  Inches(y + 0.20), Inches(1.08), Inches(0.32))
    pill.fill.solid(); pill.fill.fore_color.rgb = accent; pill.line.fill.background()
    add_text(slide, x + 4.62, y + 0.20, 1.08, 0.32, tag, 9, WHITE, True,
             PP_ALIGN.CENTER)
    slide.shapes.add_picture(str(image_path), Inches(x + 0.30), Inches(y + 0.72),
                             width=Inches(2.35), height=Inches(1.45))
    box = slide.shapes.add_textbox(Inches(x + 2.90), Inches(y + 0.76), Inches(2.85),
                                   Inches(1.46))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet; p.font.name = "Noto Sans CJK SC"; p.font.size = Pt(11.5)
        p.font.color.rgb = MUTED; p.space_after = Pt(7)
        p.level = 0; p.text = "• " + p.text


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    manifests = {
        "uji": ROOT / "data/splits/uji_pen_v2.jsonl",
        "ne": ROOT / "data/processed/natural_earth/v5.1.2_10m_coastline/manifest.jsonl",
        "usgs": ROOT / "data/processed/usgs_contours/large_scale/manifest.jsonl",
    }
    curves = {
        "synthetic": synthetic_curve(),
        "uji": representative_from_manifest(manifests["uji"]),
        "ne": representative_from_manifest(manifests["ne"]),
        "usgs": representative_from_manifest(manifests["usgs"]),
    }
    colors = {"synthetic": "#13918C", "uji": "#3068BD", "ne": "#E1802D", "usgs": "#7C52AD"}
    for key, points in curves.items():
        save_curve(points, ASSETS / f"{key}.png", colors[key])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = PAPER
    add_text(slide, 0.62, 0.36, 8.3, 0.52, "四类曲线数据集：从可控真值到真实世界泛化", 26, NAVY, True)
    add_text(slide, 0.64, 0.91, 10.4, 0.28,
             "统一任务：有序二维点序列 → 三次 B 样条拟合与内部节点选择", 12, MUTED)
    add_text(slide, 11.45, 0.44, 1.22, 0.38, "DATA", 12, TEAL, True, PP_ALIGN.CENTER)

    add_card(slide, 0.62, 1.34, "Synthetic B-spline", "有标签",
             ["在线生成；源内部节点 K = 4–56", "提供真参数、真节点与最简性证书", "用于监督训练及节点位置/数量评估"],
             ASSETS / "synthetic.png", TEAL)
    add_card(slide, 6.68, 1.34, "UJI Pen Characters v2", "真实轨迹",
             ["60 位书写者，11,640 个字符样本", "包含噪声、尖锐转向与个体书写差异", "按书写者隔离，检验跨主体泛化"],
             ASSETS / "uji.png", BLUE)
    add_card(slide, 0.62, 4.00, "Natural Earth 10m Coastline", "地理曲线",
             ["公开海岸线矢量数据，局部投影为平面曲线", "覆盖平滑弧段与多尺度海岸细节", "用于快速可视化和地理域泛化测试"],
             ASSETS / "ne.png", ORANGE)
    add_card(slide, 6.68, 4.00, "USGS 1:24,000 Contours", "地形曲线",
             ["The National Map 真实地形等高线", "局部高曲率、长曲线与区域分布变化", "用于检验复杂几何和跨区域稳健性"],
             ASSETS / "usgs.png", PURPLE)

    footer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.62), Inches(6.68),
                                    Inches(12.11), Inches(0.47))
    footer.fill.solid(); footer.fill.fore_color.rgb = NAVY; footer.line.fill.background()
    add_text(slide, 0.88, 6.68, 11.58, 0.47,
             "实验协议  |  统一重采样为 192 点 · 真实数据无节点真值 · 共同报告 MSE / 通过率 / 节点数 K / 耗时",
             11, WHITE, False, PP_ALIGN.CENTER)
    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
